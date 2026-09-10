import io
import pytest
from reportlab.pdfgen import canvas
import docx

from agents.document.agent import DocumentAgent
from agents.document.exceptions import DocumentValidationError
from agents.document.extractor import (
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_txt,
    validate_file_metadata,
)
from agents.document.providers import MockDocumentLLMProvider
from agents.document.router import (
    extract_deterministic_invoice,
    extract_deterministic_policy,
    extract_deterministic_report,
    extract_deterministic_technical_manual,
    heuristic_classify_document,
)
from agents.document.schemas import DocumentPage, DocumentType


# --- Test Helpers ---

def make_pdf(text_blocks: list[str]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for block in text_blocks:
        y = 750
        for line in block.split("\n"):
            c.drawString(100, y, line)
            y -= 20
        c.showPage()
    c.save()
    return buf.getvalue()


def make_blank_pdf(pages: int = 1) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for _ in range(pages):
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(headings: list[str], paragraphs: list[str], table_data: list[list[str]] = None) -> bytes:
    doc = docx.Document()
    for h in headings:
        doc.add_heading(h, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- 1. File Validation Tests ---

def test_valid_pdf_validation():
    pdf_bytes = make_pdf(["Test invoice content"])
    ext = validate_file_metadata("invoice.pdf", pdf_bytes, "application/pdf")
    assert ext == ".pdf"


def test_valid_docx_validation():
    docx_bytes = make_docx(["Title"], ["Content"])
    ext = validate_file_metadata("policy.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert ext == ".docx"


def test_valid_txt_validation():
    txt_bytes = b"Hello Enterprise"
    ext = validate_file_metadata("readme.txt", txt_bytes, "text/plain")
    assert ext == ".txt"


def test_unsupported_extension():
    with pytest.raises(DocumentValidationError) as exc:
        validate_file_metadata("script.py", b"print('hack')")
    assert "Unsupported file extension" in str(exc.value)


def test_invalid_mime():
    with pytest.raises(DocumentValidationError) as exc:
        validate_file_metadata("test.pdf", b"abc", mime_type="image/jpeg")
    assert "Unsupported MIME type" in str(exc.value)


def test_empty_file():
    with pytest.raises(DocumentValidationError) as exc:
        validate_file_metadata("empty.pdf", b"")
    assert "File is empty" in str(exc.value)


def test_oversized_file():
    oversized = b"a" * 1024
    with pytest.raises(DocumentValidationError) as exc:
        validate_file_metadata("huge.txt", oversized, max_size=500)
    assert "exceeds maximum permitted limit" in str(exc.value)


def test_malicious_filename():
    # Traversal in name should still extract extension cleanly
    ext = validate_file_metadata("../../../etc/passwd.pdf", b"%PDF-1.4 test")
    assert ext == ".pdf"


# --- 2. Extraction Tests ---

def test_pdf_text_extraction():
    pdf_bytes = make_pdf(["Invoice INV-2024-001", "Subtotal: $1,200.00 Total: $1,200.00"])
    pages, needs_ocr, warn = extract_text_from_pdf(pdf_bytes)
    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert "INV-2024-001" in pages[0].text
    assert needs_ocr is False


def test_pdf_ocr_required_detection():
    blank_pdf = make_blank_pdf(pages=2)
    pages, needs_ocr, warn = extract_text_from_pdf(blank_pdf)
    assert len(pages) == 2
    assert needs_ocr is True
    assert "OCR required" in warn


def test_docx_text_extraction():
    docx_bytes = make_docx(
        headings=["Section 1: Leave Entitlement"],
        paragraphs=["All full-time employees are entitled to 20 days paid leave."],
        table_data=[["Leave Type", "Days"], ["Annual", "20"], ["Sick", "10"]]
    )
    pages, sections, tables = extract_text_from_docx(docx_bytes)
    assert len(sections) == 1
    assert "Leave Entitlement" in sections[0].title
    assert len(tables) == 1
    assert tables[0].headers == ["Leave Type", "Days"]
    assert tables[0].rows[0] == ["Annual", "20"]


def test_txt_text_extraction():
    txt_content = "Line 1\r\nLine 2\r\n\r\n\r\nLine 3"
    pages = extract_text_from_txt(txt_content.encode("utf-8"))
    assert len(pages) == 1
    assert "Line 1\nLine 2\n\nLine 3" in pages[0].text


def test_txt_encodings():
    cp1252_bytes = "Café standard €50".encode("cp1252")
    pages = extract_text_from_txt(cp1252_bytes)
    assert len(pages) == 1
    assert len(pages[0].text) > 0

    latin1_bytes = "Café standard 50".encode("latin-1")
    pages_latin = extract_text_from_txt(latin1_bytes)
    assert len(pages_latin) == 1
    assert "Café" in pages_latin[0].text



def test_corrupted_pdf_handling():
    from agents.document.exceptions import DocumentExtractionError
    with pytest.raises(DocumentExtractionError):
        extract_text_from_pdf(b"not a valid pdf binary content")


# --- 3. Classification Tests ---

def test_classify_invoice():
    text = "TAX INVOICE\nVendor: Acme Corp\nBill To: Enterprise Inc\nSubtotal: $500\nTotal Amount: $550\nBalance Due"
    doc_type, conf, title = heuristic_classify_document(text, "INV-001.pdf")
    assert doc_type == DocumentType.INVOICE
    assert conf >= 0.90


def test_classify_policy():
    text = "Company Leave Policy\nEligibility: All employees\nAnnual sick leave entitlement and rules."
    doc_type, conf, title = heuristic_classify_document(text, "leave_policy.docx")
    assert doc_type == DocumentType.POLICY
    assert conf >= 0.90


def test_classify_report():
    text = "Quarterly Financial Analysis\nExecutive Summary\nKey findings indicate revenue growth of 18%."
    doc_type, conf, title = heuristic_classify_document(text, "q3_report.pdf")
    assert doc_type == DocumentType.REPORT
    assert conf >= 0.90


def test_classify_technical_manual():
    text = "Machine User Manual\nSafety Instructions\nWarning: Disconnect power before maintenance."
    doc_type, conf, title = heuristic_classify_document(text, "manual.txt")
    assert doc_type == DocumentType.TECHNICAL_MANUAL
    assert conf >= 0.90


def test_classify_unknown():
    text = "xyz 123"
    doc_type, conf, title = heuristic_classify_document(text, "misc.txt")
    assert doc_type == DocumentType.UNKNOWN
    assert conf == 0.0


# --- 4. Understanding & Domain Extraction Tests ---

@pytest.mark.asyncio
async def test_invoice_understanding_and_fields():
    invoice_text = (
        "TAX INVOICE\n"
        "Vendor: Apex Industrial Supplies\n"
        "Invoice Number: INV-98765\n"
        "Date: 2024-05-15\n"
        "Subtotal: $4,000.00\n"
        "Tax: $400.00\n"
        "Total: $4,400.00\n"
        "Hydraulic Valve  2  $2,000.00  $4,000.00"
    )
    pdf_bytes = make_pdf([invoice_text])

    agent = DocumentAgent()
    result = await agent.analyze(
        document_id="inv-001",
        file_bytes=pdf_bytes,
        filename="Apex_Invoice.pdf",
        task="extract_information"
    )

    assert result.document_type == DocumentType.INVOICE.value
    assert result.confidence >= 0.90
    assert result.structured_data.get("vendor") == "Apex Industrial Supplies"
    assert result.structured_data.get("invoice_number") == "INV-98765"
    assert result.structured_data.get("total") == "$4,400.00"
    assert len(result.sources) > 0


@pytest.mark.asyncio
async def test_policy_understanding_and_summary():
    policy_text = (
        "Enterprise Leave Policy\n"
        "Section 1: General Policy\n"
        "Eligibility: Applies to all full-time permanent employees.\n"
        "Employees must request leave at least 5 days in advance.\n"
        "Restriction: Maximum consecutive leave allowed is 15 days."
    )
    docx_bytes = make_docx(
        headings=["Enterprise Leave Policy"],
        paragraphs=[policy_text]
    )

    agent = DocumentAgent()
    result = await agent.analyze(
        document_id="pol-001",
        file_bytes=docx_bytes,
        filename="leave_policy.docx",
        task="summarize"
    )

    assert result.document_type == DocumentType.POLICY.value
    assert "leave" in result.summary.lower() or "policy" in result.summary.lower()
    assert len(result.key_points) > 0


@pytest.mark.asyncio
async def test_technical_manual_safety_extraction():
    manual_text = (
        "Operating Instructions for Milling Machine Model X\n"
        "Warning: Always wear safety goggles and protective gloves.\n"
        "Caution: Keep hands clear of moving blades during operation.\n"
        "Step 1: Press the green button to start power supply."
    )
    agent = DocumentAgent()
    result = await agent.analyze(
        document_id="man-001",
        file_bytes=manual_text.encode("utf-8"),
        filename="machine_manual.txt",
        task="extract_information"
    )

    assert result.document_type == DocumentType.TECHNICAL_MANUAL.value
    assert len(result.structured_data.get("warnings", [])) >= 2


@pytest.mark.asyncio
async def test_anti_hallucination_missing_fields():
    sparse_text = "TAX INVOICE\nSome billing text without explicit total or tax."
    pages = [DocumentPage(page_number=1, text=sparse_text)]
    inv_data, sources = extract_deterministic_invoice(sparse_text, pages)

    assert inv_data.total == "Not found in the provided document."
    assert inv_data.tax == "Not found in the provided document."


@pytest.mark.asyncio
async def test_mock_llm_provider_custom_response():
    custom_resp = {
        "document_type": "CONTRACT",
        "title": "Master Services Agreement",
        "summary": "Custom mocked legal agreement summary.",
        "key_points": ["Term: 2 years", "Governing Law: Delaware"],
        "entities": {"client": "OmniAgent"},
        "structured_data": {"liability_cap": "$1M"},
        "sources": [{"field": "liability_cap", "value": "$1M", "source": {"page": 1}}],
        "confidence": 0.99,
        "warnings": []
    }
    mock_provider = MockDocumentLLMProvider(custom_response=custom_resp)
    agent = DocumentAgent(provider=mock_provider)

    result = await agent.analyze(
        document_id="custom-01",
        file_bytes=b"Sample contract text content",
        filename="agreement.txt"
    )
    assert result.document_type == "CONTRACT"
    assert result.confidence == 0.99
    assert result.structured_data["liability_cap"] == "$1M"


@pytest.mark.asyncio
async def test_large_document_safeguard_truncation():
    large_text = "Important report findings. " + ("data " * 12000)
    assert len(large_text) > 50000

    agent = DocumentAgent()
    result = await agent.analyze(
        document_id="large-01",
        file_bytes=large_text.encode("utf-8"),
        filename="huge_report.txt",
        task="summarize"
    )
    assert result.confidence > 0.0
    assert len(result.summary) > 0
