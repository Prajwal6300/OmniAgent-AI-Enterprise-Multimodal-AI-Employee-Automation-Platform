"""
OmniAgent AI — DOCX Processing Module
Extracts text and tabular content from docx documents using python-docx.
"""

import io
from typing import Any

import structlog
from docx import Document

logger = structlog.get_logger(__name__)

DOCX_MAGIC_BYTES = b"PK\x03\x04"


def extract_docx_text(file_bytes: bytes) -> list[dict[str, Any]]:
    """
    Extracts structured text from raw DOCX bytes, including paragraphs and tables.
    """
    if not file_bytes.startswith(DOCX_MAGIC_BYTES):
        raise ValueError("Invalid DOCX file: zip/office magic bytes header mismatch")

    doc = Document(io.BytesIO(file_bytes))
    paragraphs: list[dict[str, Any]] = []

    # Extract paragraphs
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            paragraphs.append({
                "type": "paragraph",
                "index": idx + 1,
                "text": text,
            })

    # Extract tables
    for t_idx, table in enumerate(doc.tables):
        rows_data = []
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            rows_data.append(" | ".join(row_text))
        table_text = "\n".join(rows_data)
        if table_text.strip():
            paragraphs.append({
                "type": "table",
                "index": t_idx + 1,
                "text": table_text,
            })

    return paragraphs
